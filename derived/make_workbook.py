#!/usr/bin/env python3
"""Build draft/derived/TABLES.xlsx: one table per tab, grouped into the
fitness-contrast set, the stress-contrast set, and the supplementary set.

Every body is DERIVED from the v7 run tree.  Captions are PROSE and live in
CAPTIONS below; they are mirrored into draft/table_legends.md by
make_table_legends.py so the two cannot drift.

    python3 draft/derived/make_workbook.py
"""
import csv, math, re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
RUN  = ROOT / "hoffmanDrop/notAF/v7"
OUT  = HERE / "TABLES.xlsx"

COND = {"Glu": "glucose", "Ace": "acetate", "Caf": "caffeine", "NaCl": "NaCl"}
ORDER = ["Glu", "Ace", "Caf", "NaCl"]          # glucose is the control, always first


def rd(p):
    with open(p, newline="") as fh:
        return list(csv.DictReader(fh))


def ipos(x):
    """A genomic coordinate, rounded to a whole base for display."""
    return "" if x in (None, "", "NA") else int(round(float(x)))


def num(x, nd=None):
    """CSV text -> int/float where possible, else the text unchanged."""
    if x is None or x == "" or x == "NA":
        return ""
    try:
        f = float(x)
    except ValueError:
        return x
    if nd is not None:
        return round(f, nd)
    return int(f) if f == int(f) and abs(f) < 1e15 else f


# ---------------------------------------------------------------- styling ----
TITLE_F = Font(name="Arial", sz=12, b=True)
HEAD_F  = Font(name="Arial", sz=11, b=True)
BODY_F  = Font(name="Arial", sz=11)
CAP_F   = Font(name="Calibri", sz=11)
HEAD_FILL = PatternFill("solid", fgColor="FFD9D9D9")
HEAD_AL = Alignment(horizontal="center", vertical="top", wrap_text=True)
BODY_AL = Alignment(horizontal="center")
LEFT_AL = Alignment(horizontal="left")
CAP_AL  = Alignment(horizontal="left", vertical="top", wrap_text=True)


def width_for(header, values):
    w = max([len(str(header))] + [len(str(v)) for v in values] + [4])
    return min(max(w + 2, 7), 42)


def pval(v):
    """Number format for a p/q column."""
    if not isinstance(v, float):
        return None
    return "0.00E+00" if 0 < abs(v) < 0.001 else "0.0000"


def sheet(wb, tab, title, header, rows, caption, pcols=(), leftcols=(0,)):
    ws = wb.create_sheet(tab[:31])
    ncol = len(header)
    ws["A1"] = title
    ws["A1"].font = TITLE_F
    ws.row_dimensions[1].height = 16

    for j, h in enumerate(header, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font, c.fill, c.alignment = HEAD_F, HEAD_FILL, HEAD_AL

    for i, r in enumerate(rows, start=4):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_F
            c.alignment = LEFT_AL if (j - 1) in leftcols else BODY_AL
            if (j - 1) in pcols:
                nf = pval(v)
                if nf:
                    c.number_format = nf

    widths = []
    for j in range(ncol):
        w = width_for(header[j], [r[j] for r in rows])
        widths.append(w)
        ws.column_dimensions[get_column_letter(j + 1)].width = w

    # caption: merged to the exact width of the table, as many rows as it needs
    first = 4 + len(rows) + 1
    per_line = max(int(sum(widths) * 1.02), 30)
    lines = 0
    for para in caption.split("\n"):
        lines += max(1, math.ceil(len(para) / per_line))
    lines += 1                                   # slack so nothing is clipped
    last = first + lines - 1
    ws.merge_cells(start_row=first, start_column=1, end_row=last, end_column=ncol)
    cc = ws.cell(row=first, column=1, value=caption)
    cc.font, cc.alignment = CAP_F, CAP_AL
    for r in range(first, last + 1):
        ws.row_dimensions[r].height = 14

    ws.freeze_panes = "A4"
    return ws


# ------------------------------------------------------------------ inputs ---
key    = rd(RUN / "qc/results_key.csv")
loci_b = rd(RUN / "landscape/our_loci_summary_bysets.csv")
enr_b  = rd(RUN / "landscape/published_overlap_enrichment_bysets.csv")
rec_b  = rd(RUN / "landscape/published_recovery_bysets.csv")
xlev   = rd(RUN / "landscape/published_crosslevel_bysets.csv")
cover  = rd(RUN / "landscape/interval_coverage_bysets.csv")
direc  = rd(RUN / "landscape/published_crosslevel_direction.csv")
pubs   = {r["qtl"]: r for r in rd(RUN / "landscape/published_loci_used.csv")}
masked = rd(RUN / "landscape/masked_regions.csv")
val7   = rd(RUN / "candidates/fitness/qtl_candidates_fitness_validation.csv")[0]
val6   = rd(ROOT / "hoffmanDrop/notAF/v6/candidates/fitness/"
                  "qtl_candidates_fitness_validation.csv")[0]
repro  = rd(RUN / "qc/reproducibility_main.csv")
goall  = rd(RUN / "go/go_BP_all.csv")
gocand = rd(HERE / "TableS3_go_candidates.csv")
psize  = rd(HERE / "pool_size_estimates.csv")
det    = {ps: rd(RUN / f"landscape/published_crosslevel_detections_{ps}.csv")
          for ps in ("fitness", "contrast")}

main = [r for r in key if r["databin"] == "main"]

# retained peaks, and the pools and crosses they come from, recomputed here so
# the header line of every table is traceable to the peak tables themselves
MASK_REPORT = [("chrIII", 125000, 215000), ("chrIV", 15000, 60000)]


def in_mask(r):
    return any(r["seqnames"] == c and lo <= float(r["peak.pos"]) <= hi
               for c, lo, hi in MASK_REPORT)


PEAKS = {}
for ps, fn, scol in (("fitness", "fit_peaks.csv", "sample"),
                     ("contrast", "stress_peaks.csv", "L")):
    allp = [r for r in rd(RUN / "contrasts" / fn) if r["databin"] == "main"]
    kept = [r for r in allp if not in_mask(r)]
    PEAKS[ps] = dict(main=len(allp), retained=len(kept),
                     pools=len({r[scol] for r in kept}),
                     crosses=len({r["cross"] for r in kept}), rows=kept)


def median(xs):
    xs = sorted(xs)
    n = len(xs)
    if not n:
        return ""
    return xs[n // 2] if n % 2 else (xs[n // 2 - 1] + xs[n // 2]) / 2


wb = Workbook()
wb.remove(wb.active)

# ============================================================ T1  design =====
rows = []
for c in ORDER:
    w = [r for r in key if r["condition"] == c]
    m = [r for r in w if r["databin"] == "main"]
    rows.append([
        COND[c], len(w), len(m),
        sum(1 for r in w if r["databin"] == "flagged"),
        sum(1 for r in w if r["databin"] == "no_data"),
        round(median([float(r["mean_depth_seg"]) for r in m
                      if r["mean_depth_seg"] not in ("", "NA")]), 1),
        len({r["cross"] for r in m}),
    ])
rows.append(["total", len(key), len(main),
             sum(1 for r in key if r["databin"] == "flagged"),
             sum(1 for r in key if r["databin"] == "no_data"),
             "—", len({r["cross"] for r in main})])

sheet(wb, "T1 Design and yield",
      "Table 1. Wells, pools and crosses by condition",
      ["condition", "wells", "pools analysed", "flagged", "no usable estimate",
       "median depth (x)", "crosses"],
      rows,
      "Table 1. Wells, pools and crosses by condition. Of the 384 wells, 181 "
      "passed every filter and were analysed; \"pools analysed\" is the "
      "databin = main set and is the denominator for every percentage in this "
      "work. Flagged wells returned an allele-frequency estimate but carry a "
      "quality flag and are held out rather than dropped silently; wells with "
      "no usable estimate are counted separately. Depth is the median across "
      "analysed pools of mean depth over informative sites. Crosses is the "
      "number of distinct wild parents contributing at least one analysed pool "
      "in that condition, so the four condition rows sum to the 90 "
      "cross-by-condition combinations resolved here, out of the 144 designed "
      "(36 wild parents x 4 conditions); the total row gives the 35 distinct "
      "crosses that contributed anywhere. The 36th cross returned no analysable "
      "pool in any condition. Acetate was assayed for a single 48-h passage "
      "because the second timepoint was not recovered, which is why its well "
      "count converts to fewer pools.",
      leftcols=(0,))

# ================================================= T2  peak sets compared ====
cov = {r["peak_set"]: r for r in cover}
xl  = {r["peak_set"]: r for r in xlev}
dr  = {r["peak_set"]: r for r in direc}
LBL = {"fitness": "fitness contrast (allele frequency vs 0.5)",
       "contrast": "stress contrast (treated minus matched glucose)"}
rows = []
for ps in ("fitness", "contrast"):
    rows.append([
        LBL[ps],
        num(cov[ps]["peaks"]),
        PEAKS[ps]["pools"],
        PEAKS[ps]["crosses"],
        num(cov[ps]["median_kb"], 1),
        num(cov[ps]["total_interval_Mb"], 1),
        num(cov[ps]["union_Mb"], 2),
        num(cov[ps]["times_over"], 2),
        round(float(cov[ps]["union_frac"]) * 100, 1),
        num(xl[ps]["peaks_used"]),
        num(xl[ps]["obs_detections"]),
        num(xl[ps]["null_mean"], 1),
        num(xl[ps]["fold"], 2),
        float(xl[ps]["p_perm"]),
        num(dr[ps]["bbq_favoured"]),
        num(dr[ps]["pct_bbq"], 1),
    ])

sheet(wb, "T2 Peak sets compared",
      "Table 2. The two peak sets compared: detection, genome coverage and "
      "cross-level replication",
      ["peak set", "retained peaks", "pools", "crosses",
       "median interval (kb)", "total interval (Mb)", "union (Mb)",
       "coverage (x genome)", "union (% of genome)",
       "peaks in replication test", "cross-level detections", "null mean",
       "fold", "P (permutation)", "BBQ-favoured", "% BBQ-favoured"],
      rows,
      "Table 2. The two peak sets compared. Every pool yields two contrasts. "
      "The fitness contrast tests allele frequency against 0.5 and therefore "
      "detects any locus under selection, including the selection shared by "
      "every arm of the experiment. The stress contrast tests a treated pool "
      "against its matched glucose pool of the same cross and well, removing "
      "that shared component, and is the contrast comparable to Treusch et al. "
      "(2015), who used the same design. Peak counts are after MASK_REPORT "
      "(the mating-type region of chromosome III and the HO region of "
      "chromosome IV). The replication test additionally withholds the "
      "chromosome VII region under MASK_PUB, which is why it uses fewer peaks. "
      "Cross-level detections count a published locus once per cross, not once "
      "per pool, because pools of one cross share parents and segregant "
      "material and are not independent; the null is 1,000 permutations of "
      "peak position within pool, and P = 0.001 for the stress contrast is the "
      "resolution floor of that null, with 0 of 1,000 permutations reaching 35. "
      "The stress contrast recovers fewer loci in absolute terms while covering "
      "less than one genome equivalent, and replicates them better: this is the "
      "central result of the paper.",
      pcols=(13,), leftcols=(0,))

# ======================================== T3 / T8  loci by condition =========
def loci_table(ps, tab, title, caption, conds):
    src = {r["condition"]: r for r in loci_b if r["peak_set"] == ps}
    rows = []
    for c in conds:
        r = src.get(c)
        if not r:
            continue
        rows.append([COND[c], num(r["pools"]), num(r["crosses"]), num(r["peaks"]),
                     num(r["grouped_loci"]), num(r["median_kb"], 1),
                     num(r["cross_level_detections"]),
                     *([num(r["loci_overlapping_Treusch"]),
                        num(r["loci_not_in_Treusch"]),
                        num(r["their_grouped_loci"]),
                        num(r["their_cross_level"])]
                       if r["their_grouped_loci"] else ["—"] * 4)])
    sheet(wb, tab, title,
          ["condition", "pools with peaks", "crosses", "retained peaks", "grouped loci",
           "median locus width (kb)", "cross-level detections",
           "loci overlapping Treusch", "loci not in Treusch",
           "their grouped loci", "their cross-level loci"],
          rows, caption, leftcols=(0,))


loci_table("fitness", "T3 Loci by condition FIT",
           "Table 3. Peaks and grouped loci by condition, fitness contrast",
           "Table 3. Peaks and grouped loci by condition, fitness contrast. "
           "Peaks are the retained set after MASK_REPORT. Grouped loci are "
           "counted within a condition by the merging rule of Treusch et al. "
           "(2015) and are NOT additive across conditions: a locus present in "
           "several conditions is counted once in each row. Cross-level "
           "detections count each locus once per contributing cross. The last "
           "four columns exist only for caffeine and sodium chloride, the two "
           "conditions Treusch et al. assayed; glucose and acetate have no "
           "published comparator, so an em dash there means no such set exists "
           "rather than a count of zero. Read the glucose row as the internal "
           "control: it is the largest and deepest arm and its loci are the "
           "widest, which is what selection shared across every pool looks like "
           "when it is not subtracted out.",
           ORDER)

# =============================================== T4 / T9  enriched loci ======
def enr_rows(ps, unit, qmax=None, minobs=1):
    out = []
    for r in enr_b:
        if r["peak_set"] != ps or r["unit"] != unit or r["pool_set"] != "matched":
            continue
        if float(r["obs_pools"] or 0) < minobs:
            continue
        if qmax is not None and float(r["q_bh"]) >= qmax:
            continue
        p = pubs.get(r["qtl"], {})
        out.append([r["qtl"], r["qtg"] or "—", p.get("chrom", ""),
                    ipos(p.get("maxpos")), COND.get(r["pub_cond"], r["pub_cond"]),
                    num(r["pools_tested"]), num(r["obs_pools"]),
                    num(r["exp_pools"], 2), num(r["fold"], 2), float(r["q_bh"])])
    out.sort(key=lambda x: x[-1])
    return out


ENR_HEAD = ["locus", "gene", "chromosome", "position", "condition",
            "pools tested", "observed", "expected", "fold", "q"]

sheet(wb, "T4 Enriched loci FIT",
      "Table 4. Published loci enriched in the pools of their own condition, "
      "fitness contrast",
      ENR_HEAD, enr_rows("fitness", "pool", qmax=0.05),
      "Table 4. Published loci enriched in the pools of their own condition, "
      "fitness contrast. Each locus of Treusch et al. (2015) was tested against "
      "a Poisson expectation built from the fraction of the genome each pool's "
      "intervals cover, using only the pools grown in that locus's own "
      "condition, with Benjamini-Hochberg control across loci. All loci "
      "reaching q < 0.05 are shown; the complete set of 64 loci on both peak "
      "sets and both units of detection is Table S2. Two rows lie in the "
      "chromosome VII region. That region carries the growth QTL of the shared "
      "parent and this test does not withhold it, so those two rows measure our "
      "own sweep rather than replication of a published locus, and should be "
      "read as such. The unit here is the pool; counting by cross instead is "
      "the more conservative test and is given in Table S2.",
      pcols=(9,), leftcols=(0, 1, 2, 4))

sheet(wb, "T5 Negative controls FIT",
      "Table 5. Negative-control loci from conditions we never assayed, "
      "fitness contrast",
      ["locus", "gene", "chromosome", "position", "pools tested", "observed",
       "expected", "fold", "q"],
      [[r[0], r[1], r[2], r[3], r[5], r[6], r[7], r[8], r[9]]
       for r in sorted(
           ([r["qtl"], r["qtg"] or "—", pubs.get(r["qtl"], {}).get("chrom", ""),
             ipos(pubs.get(r["qtl"], {}).get("maxpos")), "",
             num(r["pools_tested"]), num(r["obs_pools"]), num(r["exp_pools"], 2),
             num(r["fold"], 2), float(r["q_bh"])]
            for r in enr_b
            if r["peak_set"] == "fitness" and r["unit"] == "pool"
            and r["pool_set"] == "all" and r["qtl"].startswith("Ehr.")),
           key=lambda x: -float(x[8]))],
      "Table 5. Negative-control loci from conditions we never assayed, fitness "
      "contrast. The four loci of Ehrenreich et al. (2010) were mapped under "
      "selections this experiment never applied, so a detection rate above "
      "chance here would indicate that the enrichment test is picking up "
      "something other than condition-specific replication. They are tested "
      "against all 181 analysed pools rather than a condition-matched subset, "
      "because no subset matches them. All four fall at or below the Poisson "
      "expectation and none approaches significance, which is the result this "
      "control is designed to produce.",
      pcols=(8,), leftcols=(0, 1, 2))

# ============================================= T6 / T10  recovery vs null ====
def rec_table(ps, tab, title, caption):
    rows = []
    for unit in ("pool", "cross"):
        for c in ("Caf", "NaCl"):
            for k in ("1", "2", "3"):
                for r in rec_b:
                    if (r["peak_set"] == ps and r["unit"] == unit
                            and r["condition"] == c and r["min_units"] == k):
                        rows.append([COND[c], unit, num(r["our_units"]),
                                     num(r["their_qtl"]), num(k),
                                     num(r["recovered"]), num(r["exp_recovered"], 1),
                                     num(r["fold"], 2), float(r["p_binom"])])
    sheet(wb, tab, title,
          ["condition", "unit of detection", "our units", "their QTL",
           "detected in at least", "recovered", "expected", "fold", "P"],
          rows, caption, pcols=(8,), leftcols=(0, 1))


rec_table("fitness", "T6 Recovery vs null FIT",
          "Table 6. Recovery of published QTL against a same-coverage null, "
          "fitness contrast",
          "Table 6. Recovery of published QTL against a same-coverage null, "
          "fitness contrast. For each condition the null is built by placing "
          "each pool's intervals at random while preserving their number and "
          "width, so a locus recovered at chance is one our intervals would "
          "have touched anyway. \"Detected in at least\" is the number of "
          "independent units required to count a locus as recovered. The unit "
          "matters: pools of one cross share parents and segregant material, so "
          "the by-pool rows overstate independence and the by-cross rows are "
          "the ones to quote. Recovery at a threshold of one is at chance in "
          "every cell, which is the saturation this peak set suffers from; the "
          "signal appears only when three independent units are required.",
          )

rec_table("contrast", "T10 Recovery vs null CON",
          "Table 10. Recovery of published QTL against a same-coverage null, "
          "stress contrast",
          "Table 10. Recovery of published QTL against a same-coverage null, "
          "stress contrast. Built exactly as Table 6 but from the stress "
          "contrast, which is the peak set comparable to Treusch et al. (2015) "
          "because they also tested a treated pool against its matched control. "
          "Fold enrichment is higher than the fitness contrast at every "
          "threshold, including a threshold of one, where the fitness contrast "
          "is at chance. The counts are smaller throughout because this peak "
          "set is smaller, so several cells rest on a handful of loci and "
          "should be read as consistent with the cross-level test of Table 2 "
          "rather than as independent evidence.",
          )

# ============================================== T7  causal gene ranks =======
LBLS = [("instances (published gene inside one of our loci)", "instances"),
        ("distinct published genes", "genes"),
        ("loci involved", "loci"),
        ("median genes per locus", "median_genes_per_locus"),
        ("mean rank percentile (0.5 = chance)", "mean_pct"),
        ("ranked first in their locus", "rank1"),
        ("expected first under a uniform-rank null", "exp_rank1"),
        ("ranked in the top five", "top5"),
        ("expected in the top five under that null", "exp_top5"),
        ("Wilcoxon P", "p_wilcox"),
        ("permutation P", "p_perm")]
rows = [[lab, num(val7.get(k, ""), 3), num(val6.get(k, ""), 3) if k in val6 else "—"]
        for lab, k in LBLS]
sheet(wb, "T7 Causal gene ranks FIT",
      "Table 7. Published causal genes ranked within our loci, fitness contrast",
      ["quantity", "min_med_depth = 2 (run of record)", "min_med_depth = 5"],
      rows,
      "Table 7. Published causal genes ranked within our loci, fitness "
      "contrast. Genes under each consensus region were ranked by the "
      "allele-series score (Materials and Methods), and the rank of each gene "
      "reported as causal by Treusch et al. (2015) was recorded. Rank "
      "percentile is the gene's rank divided by the number of genes in its "
      "locus, so 0.5 is the value expected if the score carried no information. "
      "The uniform-rank expectations are the sum over instances of 1/n and of "
      "min(5, n)/n, n being the gene count of the instance's locus. Both depth "
      "cuts are shown because min_med_depth was set to 5 a priori and relaxed "
      "to 2 after the primary result had been seen, so every primary result is "
      "reported at both. The permutation test clears no conventional threshold "
      "at either cut, so this is a descriptive result and is never described as "
      "significant. Three of the nine best-ranked instances are MSB2 inside the "
      "chromosome VII region, so the result is not independent of the decision "
      "to retain that region. The stress-contrast equivalent is not evaluated: "
      "the test is gated at ten or more published genes inside our loci and the "
      "stress contrast yields nine.",
      leftcols=(0,))

# ============================================ T8 / T9  stress contrast ======
loci_table("contrast", "T8 Loci by condition CON",
           "Table 8. Peaks and grouped loci by condition, stress contrast",
           "Table 8. Peaks and grouped loci by condition, stress contrast. The "
           "columns are those of Table 3, recomputed on the stress contrast. "
           "Glucose has no row: it is the reference each treated pool is "
           "subtracted from, so it cannot be contrasted against itself. Loci "
           "are both fewer and narrower than in the fitness contrast, and the "
           "intervals cover less than one genome equivalent in total (Table 2), "
           "which is why a published locus falling inside one of them is "
           "informative here and is not in the fitness contrast.",
           ["Ace", "Caf", "NaCl"])

sheet(wb, "T9 Enriched loci CON",
      "Table 9. Published loci enriched in the pools of their own condition, "
      "stress contrast",
      ENR_HEAD, enr_rows("contrast", "pool", qmax=0.05),
      "Table 9. Published loci enriched in the pools of their own condition, "
      "stress contrast. Built exactly as Table 4, on the peak set comparable to "
      "Treusch et al. (2015). Counting by cross rather than by pool leaves one "
      "locus at q < 0.05; that stricter unit is given for every locus in Table "
      "S2, and the cross-level test that does not depend on any single locus is "
      "in Table 2. The chromosome VII caveat of Table 4 applies here too.",
      pcols=(9,), leftcols=(0, 1, 2, 4))

# ==================================================== S1  detections ========
rows = []
for ps in ("fitness", "contrast"):
    for r in det[ps]:
        p = pubs.get(r["qtl"], {})
        rows.append([ps, r["qtl"], r["qtg"] or "—", p.get("chrom", ""),
                     ipos(p.get("maxpos")), COND.get(r["cond"], r["cond"]),
                     num(r["their_lod"], 2), num(r["their_crosses"]),
                     num(r["our_crosses"]), num(r["our_pools"])])
rows.sort(key=lambda x: (x[0] != "fitness", -float(x[8] or 0)))
sheet(wb, "S1 Cross-level detections",
      "Table S1. Cross-level detection of each published locus, both peak sets",
      ["peak set", "locus", "gene", "chromosome", "position", "condition",
       "their LOD", "their crosses", "our crosses", "our pools"],
      rows,
      "Table S1. Cross-level detection of each published locus, both peak "
      "sets. One row per locus per peak set, for the 58 loci that survive "
      "MASK_PUB and enter the replication test of Table 2. \"Our crosses\" is "
      "the statistic that test sums: a locus is counted once per cross that "
      "detects it, whatever the number of pools. \"Our pools\" is shown beside "
      "it so the difference between the two units is visible per locus, and the "
      "gap is largest exactly where one cross contributed several replicate "
      "wells. Loci with zero in both columns were tested and not detected; they "
      "are kept in the table so the denominator is explicit. Their LOD and "
      "their crosses are as published by Treusch et al. (2015) and are not "
      "recomputed here.",
      leftcols=(0, 1, 2, 3, 5))

# =============================================== S2  enrichment, all ========
rows = []
for r in enr_b:
    p = pubs.get(r["qtl"], {})
    rows.append([r["peak_set"], r["unit"], r["pool_set"], r["qtl"],
                 r["qtg"] or "—", p.get("chrom", ""), ipos(p.get("maxpos")),
                 COND.get(r["pub_cond"], r["pub_cond"] or "—"),
                 num(r["pools_tested"]), num(r["obs_pools"]),
                 num(r["exp_pools"], 3), num(r["fold"], 2),
                 float(r["p_pois"]), float(r["q_bh"])])
rows.sort(key=lambda x: (x[0] != "fitness", x[1], x[2], x[13]))
sheet(wb, "S2 Enrichment all loci",
      "Table S2. Per-locus enrichment for every published locus, both peak "
      "sets and both units of detection",
      ["peak set", "unit", "pool set", "locus", "gene", "chromosome",
       "position", "condition", "units tested", "observed", "expected",
       "fold", "P", "q"],
      rows,
      "Table S2. Per-locus enrichment for every published locus, both peak "
      "sets and both units of detection. The complete record behind Tables 4 "
      "and 9. \"Pool set\" is matched where only the pools grown in that "
      "locus's own condition were tested, and all where every analysed pool "
      "was; only the matched rows support a condition-specific claim, and the "
      "all rows exist to carry the negative controls, which have no matching "
      "condition. \"Unit\" is pool or cross. Benjamini-Hochberg control is "
      "applied within each combination of peak set, unit and pool set, so q "
      "values are comparable down a block and not across blocks. Rows with "
      "zero observed are retained so that the number of loci tested is "
      "visible.",
      pcols=(12, 13), leftcols=(0, 1, 2, 3, 4, 5, 7))

# ============================================== S3  published catalogue =====
rows = [[r["qtl"], COND.get(r["cond"], r["cond"] or "—"), r["chrom"],
         ipos(r["lo"]), ipos(r["maxpos"]), ipos(r["hi"]),
         round((int(float(r["hi"])) - int(float(r["lo"]))) / 1000, 1),
         num(r["lod"], 2), num(r["ncross"]), r["qtg"] or "—",
         r["study"], r["coord"]] for r in rd(RUN / "landscape/published_loci_used.csv")]
rows.sort(key=lambda x: (x[1], x[2], x[3]))
sheet(wb, "S3 Published loci used",
      "Table S3. The published loci this work is tested against",
      ["locus", "condition", "chromosome", "interval start", "peak position",
       "interval end", "width (kb)", "published LOD", "published crosses",
       "reported causal gene", "study", "interval definition"],
      rows,
      "Table S3. The published loci this work is tested against. Sixty of the "
      "64 are the caffeine and sodium chloride QTL of Treusch et al. (2015); "
      "the remaining four are the loci of Ehrenreich et al. (2010), mapped "
      "under selections this experiment never applied and carried as negative "
      "controls (Table 5). Coordinates are as published, on the same reference "
      "build as the gene models used here. Published crosses is the number of "
      "segregant panels in which that study detected the locus, and is the "
      "variable tested for a dose-response against our own recovery; no such "
      "relationship is present. Fifty-eight of the 64 enter the cross-level "
      "test of Table 2, the other six falling inside MASK_PUB.",
      leftcols=(0, 1, 2, 9, 10, 11))

# ================================================== S4  masked regions ======
rows = [[r["what"], r["chrom"], num(r["lo"]), num(r["hi"]),
         round((int(float(r["hi"])) - int(float(r["lo"]))) / 1000, 1),
         r["scope"], r["note"]] for r in masked]
sheet(wb, "S4 Masked regions",
      "Table S4. Regions withheld from the analysis, and where each mask applies",
      ["region", "chromosome", "start", "end", "width (kb)", "applies to", "note"],
      rows,
      "Table S4. Regions withheld from the analysis, and where each mask "
      "applies. Two masks operate. MASK_REPORT withholds the mating-type "
      "region of chromosome III and the HO region of chromosome IV, both of "
      "which are selected by the construction of the cross rather than by the "
      "condition, and applies to every peak count, every grouped locus and "
      "every figure in this work. MASK_PUB additionally withholds the "
      "chromosome VII region and applies only to the test of concordance with "
      "published loci, because that region carries the growth QTL of the "
      "shared parent and would otherwise be counted as replication of a "
      "published locus that happens to sit inside it. A peak is withheld when "
      "its summit falls inside a masked interval. Every table in this workbook "
      "states which mask it used.",
      leftcols=(0, 1, 5, 6))

# ================================================== S5  GO, best terms ======
best = {}
for r in goall:
    k = (r["peak_set"], r["condition"])
    best.setdefault(k, []).append(r)
rows = []
for ps in ("fitness", "contrast"):
    for c in ORDER:
        grp = best.get((ps, c))
        if not grp:
            continue
        grp = sorted(grp, key=lambda r: (float(r["q_perm"]), float(r["p_perm"]),
                                         -float(r["fold"])))[:10]
        for r in grp:
            rows.append([ps, COND[c], r["go"], r["term"],
                         num(r["genes_in_locus"]), num(r["term_size"]),
                         num(r["genes_tested"]), num(r["fold"], 2),
                         num(r["exp_perm"], 1), float(r["p_perm"]),
                         float(r["q_perm"]), float(r["p_hyper"])])
sheet(wb, "S5 GO best terms",
      "Table S5. The ten best-scoring GO biological-process terms per condition "
      "and peak set, against a positional null",
      ["peak set", "condition", "GO term id", "term", "genes in loci",
       "term size", "genes tested", "fold", "expected (permutation)",
       "P (permutation)", "q (permutation)", "P (hypergeometric)"],
      rows,
      "Table S5. The ten best-scoring GO biological-process terms per condition "
      "and peak set, against a positional null. Annotation is the Saccharomyces "
      "Genome Database GO association file for release R64-5-1, biological "
      "process only, propagated to ancestor terms through the is_a and part_of "
      "relations of the go-basic ontology, evidence code ND excluded. The null "
      "shifts each locus to a random genomic position 1,000 times, preserving "
      "the number and width of loci, and the permutation q is the statistic to "
      "read. Nothing reaches q < 0.1 in any condition on either peak set, so "
      "this table records a null result and names no candidate. The "
      "hypergeometric P is printed for completeness and must not be quoted: it "
      "assumes genes are exchangeable, and the genes under one locus are "
      "physical neighbours that enter or leave together, which is exactly what "
      "the positional null corrects for. On the fitness contrast the test is "
      "uninformative by construction, the glucose loci alone covering 3,556 of "
      "5,911 genes against a permuted median of 3,484.",
      pcols=(9, 10, 11), leftcols=(0, 1, 2, 3))

# =============================================== S6  GO candidate genes =====
gc_head = list(gocand[0].keys())
rows = [[num(r[k]) if k not in ("condition",) else r[k] for k in gc_head]
        for r in gocand]
sheet(wb, "S6 GO candidate genes",
      "Table S6. Genes of the expected functional class under the "
      "stress-contrast loci",
      gc_head, rows,
      "Table S6. Genes of the expected functional class under the "
      "stress-contrast loci. Genes annotated to cell wall or cell surface lying "
      "under the caffeine contrast loci, and to ion transport, ion homeostasis "
      "or osmotic response lying under the sodium chloride contrast loci. "
      "THIS IS NOT AN ENRICHMENT RESULT. The positional-null test of Table S5 "
      "returns no term at q < 0.1 in any condition on either peak set, so the "
      "proportion of genes of these classes under these loci is what a random "
      "placement of the same loci would give. Three nested flags, which must "
      "not be collapsed. in_candidate_ranking marks the genes lying inside the "
      "consensus region of a locus detected in at least two crosses, which are "
      "therefore scored at all: 24 caffeine and 13 sodium chloride. "
      "has_wild_damaging_allele marks the narrower set actually carrying at "
      "least one predicted-damaging allele in a wild parent that detects the "
      "locus: 5 and 3. nonzero_allele_series_score marks those whose composite "
      "score exceeds zero, which is a WEAKER condition: 10 and 7, because the "
      "composite also rewards proximity to the peak summit and draws its "
      "severity and conservation terms from variants on either parent, so a "
      "gene with no damaging wild allele can still score above zero. Quote the "
      "middle flag for any claim about alleles and the last one only as a "
      "ranking. The score was computed without reference to GO annotation, so "
      "it is independent of the annotation that selected these rows. No gene "
      "listed here is implicated in the trait by this experiment, none was "
      "tested by allele replacement, and the loci are wide. The table names "
      "where a follow-up might start, not what any locus does.",
      leftcols=tuple(range(len(gc_head))))

# ===================================================== S7  pool size ========
rows = [[r["sample"], r["cross"], num(r["depth"], 1), num(r["blocks"]),
         num(r["sigma"], 4), num(r["N_eff"], 0)] for r in psize]
rows.sort(key=lambda x: -float(x[2] or 0))
sheet(wb, "S7 Pool size",
      "Table S7. Effective number of segregants per pool, glucose pools",
      ["pool", "cross", "median depth (x)", "unlinked blocks",
       "sigma of allele frequency", "effective segregants"],
      rows,
      "Table S7. Effective number of segregants per pool, glucose pools. In a "
      "pool of N independent haploid segregants the allele frequency at an "
      "unselected locus departs from 0.5 with standard deviation 0.5/sqrt(N), "
      "so N follows from the spread of the smoothed trace. Sigma is the median "
      "absolute deviation of that trace sampled once per 200 kb, a spacing at "
      "which blocks are effectively unlinked, after excluding the masked "
      "regions and every peak interval called in that pool; N is 0.25 divided "
      "by its square. Across the 63 pools sequenced at 10x or better, where "
      "sigma has stopped falling with depth and is therefore a property of the "
      "pool rather than of the sequencing, the median is 229 effective "
      "segregants with an interquartile range of 136 to 283. This is a LOWER "
      "bound: residual selection, residual read noise and unequal "
      "representation all widen sigma. It also implies that the standard error "
      "the pipeline assumes, from an effective sample size of about 814 at 20x, "
      "is too small by roughly 1.9-fold, which is a quantified form of the "
      "missing overdispersion term.",
      leftcols=(0, 1))

# ============================================ S8  replicate concordance =====
rows = [[r["cross"], COND.get(r["condition"], r["condition"]),
         num(r["n_wells"]), num(r["n_intervals"]), num(r["n_shared"]),
         num(r["frac_shared"], 3)] for r in repro]
rows.sort(key=lambda x: (x[1], -float(x[5] or 0)))
sheet(wb, "S8 Replicate concordance",
      "Table S8. Concordance between replicate wells of the same cross and "
      "condition",
      ["cross", "condition", "replicate wells", "intervals", "shared",
       "fraction shared"],
      rows,
      "Table S8. Concordance between replicate wells of the same cross and "
      "condition. Where a cross was grown in more than one well of a condition, "
      "the wells are biological replicates of the same segregant pool split "
      "after sporulation, and the fraction of called intervals shared between "
      "them measures how reproducible a call is at this scale. This is the "
      "audit that justifies relaxing min_med_depth from 5 to 2: that threshold "
      "was set a priori at 5 and lowered after the primary result had been "
      "seen, so it is treated as a declared sensitivity analysis and every "
      "primary result is reported at both cuts. It is also the reason the unit "
      "of detection throughout this work is the cross and not the pool, since "
      "replicate wells of one cross are plainly not independent observations of "
      "a locus.",
      leftcols=(0, 1))

# ======================================= S9  the sequencing sample key =====
NAME = {}
for r in key:
    if r["parent_nominal"]:
        NAME[r["parent_nominal"]] = r["oldLabel_parent_a"]

rows = [[r["Sample_ID"], r["Plate"], r["Wells"], COND.get(r["condition"], r["condition"]),
         r["cross"], r["parent_final"], NAME.get(r["parent_final"], "\u2014"),
         num(r["mean_depth_seg"], 1), num(r["n_sites"]), r["databin"]] for r in key]
rows.sort(key=lambda x: (x[1], x[2]))
sheet(wb, "S9 Sample key",
      "Table S9. The sequencing sample key: every well that entered the library",
      ["sample", "plate", "well", "condition", "cross", "wild parent",
       "strain", "mean depth (x)", "informative sites", "bin"],
      rows,
      "Table S9. The sequencing sample key: every well that entered the library. "
      "All 384 wells are listed, including those later excluded, so that what was "
      "sequenced can be reconciled with what was analysed. One plate was run per "
      "condition and the same well position carries the same cross on all four, "
      "so a cross is identified by its well. The wild parent is the assignment "
      "reached from the reads (Table S10), not the plate map, and the strain "
      "column gives that parent's name in the 1011 Yeast Genomes collection where "
      "the plate map supplies one; an em dash marks the eight parents recovered by "
      "assignment that were never on the map. Depth is the mean over informative "
      "sites and informative sites is the number segregating between that cross's "
      "parents, so both depend on which parent was assigned. The bin is main for "
      "the 181 wells passing every filter, flagged for the 84 returning an "
      "estimate that carries a quality flag, and no_data for the 119 returning "
      "none; only main wells contribute to any result.",
      leftcols=(0, 1, 2, 3, 4, 5, 6, 9))

# ================================== S10  parent assignment, per well ======
seen, rows = set(), []
for r in key:
    w = r["Wells"]
    if w in seen:
        continue
    seen.add(w)
    nom, fin = r["parent_nominal"], r["parent_final"]
    rows.append([w, nom, NAME.get(nom, "\u2014"), fin, NAME.get(fin, "\u2014"),
                 "yes" if nom != fin else "\u2014",
                 r["a1_nom"], r["a1_cur"], num(r["od_final"], 2),
                 r["reassign_class"].replace("_", " "),
                 "yes" if r["unresolved"] == "TRUE" else "\u2014"])
rows.sort(key=lambda x: x[0])
sheet(wb, "S10 Parent assignment",
      "Table S10. Wild parent of each well, as arrayed and as assigned from the reads",
      ["well", "nominal parent", "nominal strain", "assigned parent",
       "assigned strain", "reassigned", "A1 fit, nominal", "A1 fit, current",
       "residual scatter", "basis", "unresolved"],
      rows,
      "Table S10. Wild parent of each well, as arrayed and as assigned from the "
      "reads. One row per well position; the assignment was identical on all four "
      "condition plates, which is itself a check, since the four plates were "
      "scored independently. The nominal parent is the plate map and the assigned "
      "parent is the strain whose genotype best explains that well's own reads "
      "under the model in Materials and Methods. Of the 96 positions, 54 were "
      "reassigned, which is 216 of the 384 wells. The two A1 columns give the "
      "verdict of the physically sequenced haploid panel on the nominal and on the "
      "then-current assignment, which is the test that settled the disagreements "
      "the 1011-panel scan could not; residual scatter is the fit statistic behind "
      "it, and values above about 1.5 mark a reassignment that was applied but "
      "flagged rather than one that is wrong, since half of all confirmed wells "
      "also sit above that value. Twelve nominal strains were confirmed at no "
      "well and eight strains absent from the plate map were recovered, so the 40 "
      "arrayed isolates resolve to 36 distinct parents. Six positions, and so 24 "
      "of the 384 wells, could not be resolved and are excluded from every analysis.",
      leftcols=(0, 1, 2, 3, 4, 5, 6, 7, 9, 10))

ORDER_TABS = ["T1 Design and yield", "T2 Peak sets compared",
              "T3 Loci by condition FIT", "T4 Enriched loci FIT",
              "T5 Negative controls FIT", "T6 Recovery vs null FIT",
              "T7 Causal gene ranks FIT", "T8 Loci by condition CON",
              "T9 Enriched loci CON", "T10 Recovery vs null CON",
              "S1 Cross-level detections", "S2 Enrichment all loci",
              "S3 Published loci used", "S4 Masked regions",
              "S5 GO best terms", "S6 GO candidate genes",
              "S7 Pool size", "S8 Replicate concordance",
              "S9 Sample key", "S10 Parent assignment"]
assert sorted(ORDER_TABS) == sorted(ws.title for ws in wb), "tab list out of step"
wb._sheets = [wb[t] for t in ORDER_TABS]
wb.active = 0

wb.save(OUT)
print(f"wrote {OUT}")
for ws in wb:
    print(f"  {ws.title:32s} {ws.max_row - 1:5d} rows x {ws.max_column} cols")
