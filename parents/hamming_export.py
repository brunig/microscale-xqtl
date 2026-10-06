#!/usr/bin/env python3
"""
Extract clean CSVs from the Hamming distance Excel, preserving the
color-coded confidence flags on cell fills.

Reads:  ~/Desktop/AFxQTL/tables/hamming_resultsLongerShuffle_relabelled.xlsx
Writes: /tmp/hamming_reference_haploids.csv
        /tmp/hamming_cross_relabels.csv

Reference haploids: rows with STRAIN != EMPTY. These map a common
strain name (short96Name) to a Hamming-verified 1011 code
(postHam / Standardized_name). Used to validate xQTLpaper labels.

Cross relabels: rows with STRAIN == EMPTY. These are the cross-pool
samples from a previous run where old label != Hamming call.

Confidence:
    green   = fill FFC6EFCE  (clean top-1 Hamming hit)
    yellow  = fill FFFFEB9C  (ambiguous, alternates in parent_a_alt)
    none    = no fill        (either reference haploid or no relabel)
"""
import csv
import os
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, '-m', 'pip', 'install',
                           '--user', '--quiet', 'openpyxl'])
    import openpyxl

XLSX = str(Path.home() / 'Desktop/AFxQTL/tables/hamming_resultsLongerShuffle_relabelled.xlsx')
OUT_REF   = '/tmp/hamming_reference_haploids.csv'
OUT_CROSS = '/tmp/hamming_cross_relabels.csv'

GREEN  = 'FFC6EFCE'
YELLOW = 'FFFFEB9C'

wb = openpyxl.load_workbook(XLSX, data_only=False)
ws = wb['hamming_resultsLonger']

headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

def get(r, name):
    c = headers.get(name)
    return ws.cell(r, c).value if c else None

def fill_color(r, name):
    c = headers.get(name)
    if not c:
        return None
    cell = ws.cell(r, c)
    if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb:
        rgb = cell.fill.fgColor.rgb
        if rgb == GREEN:  return 'green'
        if rgb == YELLOW: return 'yellow'
    return None

# ---- Reference haploids ------------------------------------------
# short96Name = common strain name (e.g. CLIB192)
# postHam / Standardized_name = Hamming-verified 1011 code
# col1Value / col2Value = top-1 and top-2 Hamming distances
ref_rows = []
for r in range(2, ws.max_row + 1):
    strain = get(r, 'STRAIN')
    if strain in (None, '', 'EMPTY'):
        continue
    ref_rows.append({
        'Sample_ID':          get(r, 'Sample_ID'),
        'plate_location':     get(r, 'LOCATION_BoxA1'),
        'strain_full':        strain,
        'strain_common_name': get(r, 'short96Name'),
        'postHam':            get(r, 'postHam'),
        'standardized_name':  get(r, 'Standardized_name'),
        'top1_code':          get(r, 'col1Name'),
        'top1_hamming':       get(r, 'col1Value'),
        'top2_code':          get(r, 'col2Name'),
        'top2_hamming':       get(r, 'col2Value'),
        'top3_code':          get(r, 'col3Name'),
        'top3_hamming':       get(r, 'col3Value'),
    })

with open(OUT_REF, 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=ref_rows[0].keys())
    w.writeheader()
    w.writerows(ref_rows)
print(f'Wrote {len(ref_rows)} reference haploid rows -> {OUT_REF}')


# ---- Cross-pool relabels ------------------------------------------
# STRAIN == EMPTY, but parent_a is filled with the corrected code.
# Confidence comes from the cell fill on parent_a.
cross_rows = []
for r in range(2, ws.max_row + 1):
    strain = get(r, 'STRAIN')
    if strain != 'EMPTY':
        continue
    conf = fill_color(r, 'parent_a')
    cross_rows.append({
        'Sample_ID':          get(r, 'Sample_ID'),
        'Plate':              get(r, 'Plate'),
        'Wells':              get(r, 'Wells'),
        'parent_alpha':       get(r, 'parent_alpha'),
        'parent_a_oldLabel':  get(r, 'parent_a_oldLabel'),
        'parent_a':           get(r, 'parent_a'),
        'parent_a_alt':       get(r, 'parent_a_alt'),
        'confidence':         conf or 'none',
        'Drug3Let':           get(r, 'Drug3Let'),
    })

with open(OUT_CROSS, 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=cross_rows[0].keys())
    w.writeheader()
    w.writerows(cross_rows)
print(f'Wrote {len(cross_rows)} cross-relabel rows -> {OUT_CROSS}')


# ---- Summary ------------------------------------------------------
print('\nReference haploids by strain (unique common names):')
strain_counts = {}
for row in ref_rows:
    n = row['strain_common_name']
    strain_counts[n] = strain_counts.get(n, 0) + 1
print(f'  Unique strains: {len(strain_counts)}')

print('\nCross relabels by confidence:')
conf_counts = {}
for row in cross_rows:
    conf_counts[row['confidence']] = conf_counts.get(row['confidence'], 0) + 1
for k, v in sorted(conf_counts.items()):
    print(f'  {k:10s}: {v}')

print('\nMismatches (parent_a_oldLabel != parent_a) in cross rows:')
n_mismatch = 0
for row in cross_rows:
    old = row['parent_a_oldLabel']
    new = row['parent_a']
    if old and new and old != new and old != 'NA' and new != 'NA':
        n_mismatch += 1
print(f'  Total: {n_mismatch}')
