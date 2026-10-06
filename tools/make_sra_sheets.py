#!/usr/bin/env python3
"""Build the NCBI BioSample and SRA metadata sheets for the raw reads.

ONE submission holding both sets:
  pools    384 segregant pools (plate map: parents/flatkey_A1.csv)
  parents  the haploid reference panel, plate A1ref (non-EMPTY wells)
The two sets reuse FASTQ names (Sample001_S1_...), so parent files are
uploaded under the prefix PARENT_PREFIX and named that way in the SRA sheet.
Outputs: sra/biosample_all.xlsx and sra/sra_metadata_all.tsv (+ per-set TSVs).

BioSample package: Microbe 1.0 (NCBI: unicellular microbes where MIxS does
not fit; lab pools have no environment fields). Template: sra/template_Microbe.1.0.xlsx
and the filled copies are sra/biosample_*.xlsx, ready to upload. Attributes describe the biological sample
only; analysis status (main / flagged) belongs in the paper, not here, and
changes between runs.

If sra/fastq_sizes.tsv exists (from tools/sra_inventory.sh), every
filename written to the SRA sheet is checked against it, so a sheet never
names a file that is not on Hoffman2.

    python3 tools/make_sra_sheets.py
"""
import csv, os, re, sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
OUT = os.path.join(REPO, 'sra')
POOL_KEY = os.path.join(REPO, 'parents', 'flatkey_A1.csv')
PARENT_KEY = os.path.join(REPO, 'parents', 'refA1key.csv')
INVENTORY = os.path.join(OUT, 'fastq_sizes.tsv')   # from tools/sra_inventory.sh

COND = {'Glu': 'glucose (control)', 'Ace': 'acetic acid', 'Caf': 'caffeine',
        'NaCl': 'sodium chloride'}
REGIME = {'Glu': 'two 48-h passages in HCD medium',
          'Caf': 'two 48-h passages in HCD medium',
          'NaCl': 'two 48-h passages in HCD medium',
          'Ace': 'one 48-h passage in HCD medium'}
PARENT_PREFIX = 'A1ref_'                # parent FASTQs renamed on upload
NA_LAB = 'missing: lab stock'          # INSDC null term for laboratory material
INSTRUMENT = 'Illumina NovaSeq X'
DESIGN = ('Genomic DNA from a pool of haploid segregants (or a haploid parent '
          'culture), Tn5 tagmentation (homebrew, after Picelli et al. 2014), '
          'dual-indexed PCR with IDT for Illumina Nextera DNA UD indexes, '
          '2 x 150 bp paired-end sequencing.')

TEMPLATE = os.path.join(OUT, 'template_Microbe.1.0.xlsx')
# custom attributes (BioSample accepts extra columns) appended after the template's
CUSTOM_POOLS = ['treatment', 'growth_protocol', 'cross', 'wild_parent', 'common_parent', 'plate_map_wild_parent',
                'parent_assignment', 'plate', 'well']
CUSTOM_PARENTS = ['plate', 'wells', 'yeast_1011_genomes_code']
SRA_COLS = ['sample_name', 'library_ID', 'title', 'library_strategy',
            'library_source', 'library_selection', 'library_layout',
            'platform', 'instrument_model', 'design_description', 'filetype',
            'filename', 'filename2', 'filename3', 'filename4']


def template_cols():
    import openpyxl
    ws = openpyxl.load_workbook(TEMPLATE).active
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == '*sample_name':
            return [c for c in row if c]
    sys.exit('no header row in template')


def write_biosample(stem, rows, custom):
    """TSV plus the NCBI template filled in place (comment rows kept)."""
    import openpyxl
    cols = template_cols() + custom
    unknown = {k for r in rows for k in r} - set(cols)
    if unknown:
        sys.exit(f'attributes not in template or custom list: {unknown}')
    write(os.path.join(OUT, stem + '.tsv'), cols, rows)
    wb = openpyxl.load_workbook(TEMPLATE); ws = wb.active
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), 1)
               if row and row[0] == '*sample_name')
    for j, c in enumerate(cols, 1):
        ws.cell(hdr, j, c)
    for i, r in enumerate(rows, hdr + 1):
        for j, c in enumerate(cols, 1):
            ws.cell(i, j, r.get(c, ''))
    wb.save(os.path.join(OUT, stem + '.xlsx'))
    print(f'{len(rows):4d} rows  sra/{stem}.xlsx')


def write(path, cols, rows):
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, cols, delimiter='\t', lineterminator='\n', extrasaction='raise', restval='')
        w.writeheader()
        w.writerows(rows)
    print(f'{len(rows):4d} rows  {os.path.relpath(path, REPO)}')


def inventory():
    """{set: {basename: path}} from sra_inventory.sh, or None."""
    if not os.path.exists(INVENTORY):
        return None
    inv = defaultdict(dict)
    for r in csv.DictReader(open(INVENTORY), delimiter='\t'):
        inv[r['set']][os.path.basename(r['path'])] = r['path']
    return inv


def lanes_for(inv, sset, sid, s):
    """FASTQ basenames for one library: every lane, R1 then R2 per lane."""
    pat = re.compile(rf'^{sid}_S{s}_L00(\d)_R([12])_001\.fastq\.gz$')
    if inv is None:                      # before the inventory: assume 2 lanes
        return [f'{sid}_S{s}_L00{l}_R{r}_001.fastq.gz' for l in (1, 2) for r in (1, 2)]
    hits = sorted((m.group(1), m.group(2), b) for b in inv[sset] if (m := pat.match(b)))
    return [b for _, _, b in hits]


def main():
    os.makedirs(OUT, exist_ok=True)
    inv = inventory()
    problems = []

    ## pools
    pools = list(csv.DictReader(open(POOL_KEY)))
    bs, sra = [], []
    for i, r in enumerate(pools, start=1):          # S-number = key row order
        sid, cond = r['Sample_ID'], r['Drug3Let']
        wild = r['parent_final']
        name = f'xQTL_{sid}_{wild}xBBQ_{cond}'
        unresolved = r['unresolved'] == 'TRUE'
        bs.append({
            '*sample_name': name,
            'sample_title': f'Segregant pool, {wild} x BBQ, {COND[cond]}, well {r["Well_Position"]}',
            'bioproject_accession': '',
            '*organism': 'Saccharomyces cerevisiae',
            'strain': '',
            'isolate': f'{wild}xBBQ segregant pool {sid}',
            'isolation_source': f'laboratory cross {wild} x BBQ; haploid segregants selected in {COND[cond]} ({REGIME[cond]})',
            '*collection_date': NA_LAB,
            '*geo_loc_name': NA_LAB,
            '*sample_type': 'pooled haploid segregant culture',
            'genotype': f'haploid segregants of {wild} x BBQ',
            'treatment': COND[cond],
            'growth_protocol': REGIME[cond],
            'cross': f'{wild} x BBQ', 'wild_parent': wild, 'common_parent': 'BBQ',
            'plate_map_wild_parent': r['parent_nominal'],
            'parent_assignment': 'unresolved' if unresolved else (
                'reassigned from sequence' if wild != r['parent_nominal'] else 'confirmed from sequence'),
            'plate': r['Plate'], 'well': r['Wells'],
            'description': (
                f'Wild parent not resolved from sequence; best call {wild}, plate map named {r["parent_nominal"]}.'
                if unresolved else
                f'Wild parent identified from sequence; plate map named {r["parent_nominal"]}.'
                if wild != r['parent_nominal'] else
                'Wild parent as on the plate map, confirmed from sequence.'),
        })
        files = lanes_for(inv, 'pools', sid, i)
        if len(files) != 4:
            problems.append(f'pools {sid}: {len(files)} files')
        sra.append({
            'sample_name': name, 'library_ID': f'xQTL_{sid}',
            'title': f'WGS of S. cerevisiae segregant pool ({wild} x BBQ, {COND[cond]})',
            'library_strategy': 'WGS', 'library_source': 'GENOMIC',
            'library_selection': 'RANDOM', 'library_layout': 'paired',
            'platform': 'ILLUMINA', 'instrument_model': INSTRUMENT,
            'design_description': DESIGN, 'filetype': 'fastq',
            **{c: (files[k] if k < len(files) else '')
               for k, c in enumerate(['filename', 'filename2', 'filename3', 'filename4'])},
        })
    write_biosample('biosample_pools', bs, CUSTOM_POOLS)
    write(os.path.join(OUT, 'sra_metadata_pools.tsv'), SRA_COLS, sra)
    all_bs, all_sra = list(bs), list(sra)

    ## parents
    if not os.path.exists(PARENT_KEY):
        sys.exit(f'missing {PARENT_KEY}')
    par = list(csv.DictReader(open(PARENT_KEY, encoding='utf-8-sig')))
    # One BioSample per STRAIN; a strain sequenced in two wells is two
    # libraries (SRA rows) of one sample. NCBI rejects BioSamples that differ
    # only in name, which per-well samples of the same strain would.
    bs, sra, seen = [], [], {}
    for i, r in enumerate(par, start=1):
        if not r['Project'] or r['STRAIN'] == 'EMPTY':
            continue
        sid, strain, code = r['Sample_ID'], r['STRAIN'], r['Standardized_name']
        name = f'A1ref_{strain}'
        if name not in seen:
            seen[name] = []
            bs.append({
                '*sample_name': name,
                'sample_title': f'Haploid reference strain {strain}',
                'bioproject_accession': '',
                '*organism': 'Saccharomyces cerevisiae',
                'strain': strain, 'isolate': '',
                'isolation_source': 'laboratory haploid derivative of a wild isolate',
                '*collection_date': NA_LAB, '*geo_loc_name': NA_LAB,
                '*sample_type': 'haploid strain culture',
                'genotype': 'haploid',
                'plate': 'A1ref', 'wells': '', 'yeast_1011_genomes_code': code,
                'description': (f'1011 Yeast Genomes code {code}. ' if code else '')
                               + 'Sequenced as the reference panel for assigning wild parents to pools.',
            })
        seen[name].append(r['Wells'])
        files = [PARENT_PREFIX + f for f in lanes_for(inv, 'parents', sid, i)]
        if inv is not None and len(files) not in (2, 4):
            problems.append(f'parents {sid}: {len(files)} files')
        sra.append({
            'sample_name': name, 'library_ID': f'A1ref_{sid}',
            'title': f'WGS of S. cerevisiae haploid strain {strain} (plate A1ref well {r["Wells"]})',
            'library_strategy': 'WGS', 'library_source': 'GENOMIC',
            'library_selection': 'RANDOM', 'library_layout': 'paired',
            'platform': 'ILLUMINA', 'instrument_model': INSTRUMENT,
            'design_description': DESIGN, 'filetype': 'fastq',
            **{c: (files[k] if k < len(files) else '')
               for k, c in enumerate(['filename', 'filename2', 'filename3', 'filename4'])},
        })
    for b in bs:
        w = seen[b['*sample_name']]
        b['wells'] = ', '.join(w)
        b['description'] += f' Plate A1ref well{"s" if len(w) > 1 else ""} {", ".join(w)}.'
    write_biosample('biosample_parents', bs, CUSTOM_PARENTS)
    write(os.path.join(OUT, 'sra_metadata_parents.tsv'), SRA_COLS, sra)

    ## the combined submission
    write_biosample('biosample_all', all_bs + bs, CUSTOM_POOLS + [c for c in CUSTOM_PARENTS if c not in CUSTOM_POOLS])
    write(os.path.join(OUT, 'sra_metadata_all.tsv'), SRA_COLS, all_sra + sra)
    fn = [r[c] for r in all_sra + sra for c in ('filename', 'filename2', 'filename3', 'filename4') if r[c]]
    if len(fn) != len(set(fn)):
        problems.append('duplicate filenames in the combined SRA sheet')

    if inv is None:
        print('NOTE: no sra/fastq_sizes.tsv yet; pool filenames assume 2 lanes, '
              'parent filenames are UNCHECKED. Re-run after sra_inventory.sh.')
    for p in problems:
        print('PROBLEM', p)


if __name__ == '__main__':
    main()
